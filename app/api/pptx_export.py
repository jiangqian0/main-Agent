from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Optional
import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from app.service.workspace_service import WorkspaceService
import os
import uuid

router = APIRouter()


class ExportRequest(BaseModel):
    html_path: str
    output_name: Optional[str] = None


def hex_to_rgb(hex_color: str) -> tuple:
    """Convert hex color to RGB tuple"""
    hex_color = hex_color.lstrip('#')
    if len(hex_color) == 6:
        return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
    return (0, 0, 0)


def extract_color_from_css(css_text: str, var_name: str) -> Optional[str]:
    """Extract color value from CSS variables"""
    pattern = rf'--{var_name}:\s*([^;]+);'
    match = re.search(pattern, css_text)
    if match:
        return match.group(1).strip()
    return None


def parse_html_to_slides(html_content: str) -> list:
    """Parse HTML content into slide data"""
    slides = []

    # Extract CSS variables for theme colors
    ink_color = extract_color_from_css(html_content, 'ink') or '#0a0a0b'
    paper_color = extract_color_from_css(html_content, 'paper') or '#f1efea'

    # Find all slide sections
    slide_pattern = r'<section[^>]*class="slide[^"]*"[^>]*>(.*?)</section>'
    slide_matches = re.findall(slide_pattern, html_content, re.DOTALL)

    for idx, slide_content in enumerate(slide_matches):
        slide_data = {
            'title': '',
            'content': [],
            'is_dark': 'dark' in slide_content,
            'index': idx + 1
        }

        # Extract title classes
        title_patterns = [
            r'<h1[^>]*class="[^"]*h-hero[^"]*"[^>]*>(.*?)</h1>',
            r'<h1[^>]*class="[^"]*display[^"]*"[^>]*>(.*?)</h1>',
            r'<h2[^>]*class="[^"]*h-hero[^"]*"[^>]*>(.*?)</h2>',
            r'<h1[^>]*>(.*?)</h1>',
            r'<h2[^>]*>(.*?)</h2>',
        ]

        for pattern in title_patterns:
            title_match = re.search(pattern, slide_content, re.DOTALL)
            if title_match:
                slide_data['title'] = re.sub(r'<[^>]+>', '', title_match.group(1)).strip()
                break

        # Extract body content
        content_patterns = [
            r'<p[^>]*class="[^"]*lead[^"]*"[^>]*>(.*?)</p>',
            r'<p[^>]*class="[^"]*body[^"]*"[^>]*>(.*?)</p>',
            r'<p[^>]*>(.*?)</p>',
        ]

        for pattern in content_patterns:
            content_matches = re.findall(pattern, slide_content, re.DOTALL)
            for match in content_matches:
                text = re.sub(r'<[^>]+>', '', match).strip()
                if text and len(text) > 5:
                    slide_data['content'].append(text)

        # Extract list items
        li_pattern = r'<li[^>]*>(.*?)</li>'
        li_matches = re.findall(li_pattern, slide_content, re.DOTALL)
        for match in li_matches:
            text = re.sub(r'<[^>]+>', '', match).strip()
            if text:
                slide_data['content'].append(f"• {text}")

        # Extract stat cards
        stat_pattern = r'<div[^>]*class="[^"]*stat-card[^"]*"[^>]*>(.*?)</div>'
        stat_matches = re.findall(stat_pattern, slide_content, re.DOTALL)
        for match in stat_matches:
            # Extract stat number
            num_match = re.search(r'stat-nb[^>]*>(.*?)</', match, re.DOTALL)
            label_match = re.search(r'stat-label[^>]*>(.*?)</', match, re.DOTALL)
            if num_match:
                num = re.sub(r'<[^>]+>', '', num_match.group(1)).strip()
                label = re.sub(r'<[^>]+>', '', label_match.group(1)).strip() if label_match else ''
                slide_data['content'].append(f"{num} {label}")

        if slide_data['title'] or slide_data['content']:
            slides.append(slide_data)

    return slides


@router.post("/export-pptx")
async def export_html_to_pptx(request: ExportRequest):
    """Export HTML PPT to PPTX format"""
    try:
        service = WorkspaceService()

        # Read the HTML file
        html_content = service.read_file(request.html_path)
        if html_content is None:
            raise HTTPException(status_code=404, detail="HTML file not found")

        # Parse HTML to slides
        slides_data = parse_html_to_slides(html_content)

        if not slides_data:
            raise HTTPException(status_code=400, detail="No slides found in HTML file")

        # Create PPTX
        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9 aspect ratio
        prs.slide_height = Inches(7.5)

        # Extract theme colors
        ink_color = hex_to_rgb(extract_color_from_css(html_content, 'ink') or '#0a0a0b')
        paper_color = hex_to_rgb(extract_color_from_css(html_content, 'paper') or '#f1efea')

        for slide_data in slides_data:
            # Use blank layout
            slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(slide_layout)

            # Set background color
            if slide_data['is_dark']:
                background = slide.background
                fill = background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(*ink_color)
            else:
                background = slide.background
                fill = background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(*paper_color)

            # Add title
            if slide_data['title']:
                left = Inches(0.5)
                top = Inches(0.3)
                width = Inches(12.333)
                height = Inches(1.2)
                title_box = slide.shapes.add_textbox(left, top, width, height)
                tf = title_box.text_frame
                tf.word_wrap = True
                tf.text = slide_data['title']
                p = tf.paragraphs[0]
                p.font.size = Pt(40)
                p.font.bold = True
                if slide_data['is_dark']:
                    p.font.color.rgb = RGBColor(*paper_color)
                else:
                    p.font.color.rgb = RGBColor(*ink_color)

            # Add content
            if slide_data['content']:
                left = Inches(0.5)
                top = Inches(1.8)
                width = Inches(12.333)
                height = Inches(5.2)
                content_box = slide.shapes.add_textbox(left, top, width, height)
                tf = content_box.text_frame
                tf.word_wrap = True

                for i, text in enumerate(slide_data['content']):
                    if i == 0:
                        p = tf.paragraphs[0]
                    else:
                        p = tf.add_paragraph()
                    p.text = text
                    p.font.size = Pt(20)
                    p.line_spacing = 1.5
                    if slide_data['is_dark']:
                        p.font.color.rgb = RGBColor(*paper_color)
                    else:
                        p.font.color.rgb = RGBColor(*ink_color)

            # Add page number
            page_num_box = slide.shapes.add_textbox(
                Inches(12.5), Inches(7), Inches(0.7), Inches(0.4)
            )
            tf = page_num_box.text_frame
            tf.text = str(slide_data['index'])
            p = tf.paragraphs[0]
            p.font.size = Pt(12)
            p.alignment = PP_ALIGN.RIGHT
            if slide_data['is_dark']:
                p.font.color.rgb = RgbColor(*paper_color)
            else:
                p.font.color.rgb = RgbColor(*ink_color)

        # Generate output filename
        if request.output_name:
            output_filename = request.output_name
            if not output_filename.endswith('.pptx'):
                output_filename += '.pptx'
        else:
            html_name = Path(request.html_path).stem
            output_filename = f"{html_name}.pptx"

        # Save to temp file first
        temp_dir = Path("app/data/exports")
        temp_dir.mkdir(parents=True, exist_ok=True)
        temp_filename = f"{uuid.uuid4().hex[:8]}_{output_filename}"
        temp_path = temp_dir / temp_filename

        prs.save(str(temp_path))

        return {
            "success": True,
            "message": "PPTX exported successfully",
            "download_path": f"/api/workspace/download/{temp_filename}",
            "filename": output_filename,
            "slide_count": len(slides_data)
        }

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/download/{filename}")
async def download_pptx(filename: str):
    """Download the generated PPTX file"""
    file_path = Path("app/data/exports") / filename

    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")

    # Remove the UUID prefix from filename
    parts = filename.split('_', 1)
    original_name = parts[1] if len(parts) > 1 else filename

    return FileResponse(
        str(file_path),
        media_type='application/vnd.openxmlformats-officedocument.presentationml.presentation',
        filename=original_name,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{original_name}"}
    )


@router.get("/ppt-files")
async def list_ppt_files():
    """List HTML files in workspace that can be exported"""
    service = WorkspaceService()
    files = service.list_files()

    ppt_files = []
    for f in files:
        if f['name'].endswith('.html') and 'index' in f['name'].lower():
            ppt_files.append(f)

    return ppt_files
